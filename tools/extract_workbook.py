import openpyxl, sys, json, datetime
p, out = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(p, data_only=True, read_only=True)
ws = wb["Trip Log"]
cols = ["dateDispatched","dateDelivered","plate","assetCode","vehicleType","type","costCenter","contractType","client","loadingFrom","loadingTo","journeyPlanId","waybill","returnWaybill","invoiceStatus","sellingPrice","vat","billingAfterVat","quoteNo","soNo","manualInvoiceNo","invoiceNo","paymentStatus","advance","amountPaid","balance","total","month","eligible","rank","ae","note","ag","ah"]
def conv(v):
    if isinstance(v, datetime.datetime): return v.strftime("%Y-%m-%d")
    if isinstance(v, str): 
        v=v.strip()
        return v if v!="" else None
    return v
trips=[]
for i,row in enumerate(ws.iter_rows(min_row=2, max_row=3709, max_col=34, values_only=True), start=2):
    d=dict(zip(cols,[conv(v) for v in row]))
    if d["client"] is None and d["plate"] is None and d["sellingPrice"] is None: continue
    d["xlRow"]=i
    trips.append(d)
print("trips:", len(trips))
# assets
ws2=wb["Asset List"]; assets=[]
for row in ws2.iter_rows(min_row=2, values_only=True):
    if row[0]: assets.append(dict(zip(["plate","assetCode","vehicleType","costCenter","type"],[conv(v) for v in row[:5]])))
ws3=wb["Settings"]; clients=[]
for row in ws3.iter_rows(min_row=58, max_row=113, max_col=1, values_only=True):
    if row[0] and str(row[0]).strip(): clients.append(str(row[0]).strip())
# cached summaries for validation
ws4=wb["Client Summary"]; cs=[]
for row in ws4.iter_rows(min_row=8, max_row=63, max_col=5, values_only=True):
    if row[0]: cs.append(list(row))
ws5=wb["Statement CLient "]; st=[]
for row in ws5.iter_rows(min_row=8, max_row=63, max_col=7, values_only=True):
    if row[0]: st.append(list(row))
json.dump({"trips":trips,"assets":assets,"clients":clients,"clientSummary":cs,"statement":st}, open(out,"w"), default=str)
# quick validation of formulas against cached values
bad=0
for t in trips:
    p=t["sellingPrice"]
    if isinstance(p,(int,float)):
        vat=round(p*0.075+1e-9,2)
        if t["vat"] is not None and abs(vat-float(t["vat"]))>0.011: bad+=1
print("vat mismatches:", bad)
import collections
print("sum selling:", sum(t["sellingPrice"] for t in trips if isinstance(t["sellingPrice"],(int,float))))
print("sum billing:", sum(t["billingAfterVat"] for t in trips if isinstance(t["billingAfterVat"],(int,float))))
print("sum paid:", sum(t["amountPaid"] for t in trips if isinstance(t["amountPaid"],(int,float))))
print("sum balance:", sum(t["balance"] for t in trips if isinstance(t["balance"],(int,float))))
print("nonnumeric selling:", [ (t["xlRow"],t["sellingPrice"]) for t in trips if not isinstance(t["sellingPrice"],(int,float))][:20])
print("statement sample:", st[:3])
print("clients:", len(clients))
# duplicates check
keys=collections.Counter((t["dateDispatched"],t["plate"],t["client"],t["waybill"],t["journeyPlanId"]) for t in trips)
print("dup keys:", sum(1 for k,c in keys.items() if c>1), [k for k,c in keys.items() if c>1][:10])
