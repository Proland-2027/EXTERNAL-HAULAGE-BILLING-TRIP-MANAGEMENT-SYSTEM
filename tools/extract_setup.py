#!/usr/bin/env python3
"""Extracts every sheet of the Setup_Haulage workbook (assets, items, staff, suppliers, LOB, rate card, KPI, expense
accounts, All Routes, driver mapping, KM tables, zones, 3PL rate cards, spelling corrections) into one JSON file that
tools/build_setup_masters.js turns into the application masters. Cached values only; the workbook is never modified."""
import json, sys, datetime
from openpyxl import load_workbook

src = sys.argv[1] if len(sys.argv) > 1 else 'data/workbook/Setup_Haulage_original.xlsx'
out = sys.argv[2] if len(sys.argv) > 2 else 'data/workbook/setup_extract.json'
wb = load_workbook(src, data_only=True)

def clean(v):
    if isinstance(v, datetime.datetime): return v.date().isoformat()
    if isinstance(v, datetime.date): return v.isoformat()
    if isinstance(v, str): return v.strip()
    return v

def table(name, header_row=1, key_col=0):
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    headers = [clean(h) or f'col{i}' for i, h in enumerate(rows[header_row - 1])]
    out = []
    for r in rows[header_row:]:
        vals = [clean(v) for v in r]
        if all(v in (None, '') for v in vals): continue
        if key_col is not None and vals[key_col] in (None, ''): continue
        out.append({headers[i]: vals[i] for i in range(len(headers)) if headers[i]})
    return {'headers': headers, 'rows': out}

data = {
    'source': src.split('/')[-1], 'extractedAt': datetime.datetime.utcnow().replace(microsecond=0).isoformat() + 'Z',
    'assets': table('Asset Code'),
    'items': table('Item List'),
    'staff': table('Staffs List', key_col=1),
    'suppliers': table('Supplier', header_row=4),
    'lob': [clean(r[1]) for r in wb['LOB'].iter_rows(values_only=True) if clean(r[1])],
    'ratecard': table('Rate Card for costing ', header_row=2, key_col=3),
    'kpi': [[clean(v) for v in r] for r in wb['KPI'].iter_rows(values_only=True) if any(clean(v) not in (None, '') for v in r)],
    'accounts': table('Expense Account ', key_col=None),
    'allRoutes': table('All Routes'),
    'driverMapping': table('Driver Mapping On Each Vehicle '),
    'kmTables': {n: table(n) for n in ['KM from Lagos', 'KM from Oyo (Ibadan)', 'KM from Rivers (PH)', 'KM from FCT (Abuja)', 'KM from Nasarawa (Keffi)']},
    'internalZones': table('Internal Zones'),
    'zoneAreas': table('Zone Areas'),
    'tplRateInternal': table('3PL Rate Card - Internal'),
    'tplRateInterstate': table('3PL Rate Card - Interstate'),
    'spelling': table('Spelling Corrections'),
    'method': [[clean(v) for v in r] for r in wb['Method & Notes'].iter_rows(values_only=True) if clean(r[0])],
}
# the supplier sheet carries two header rows above the column names: scope and expense type per equipment column
ws = wb['Supplier']
data['suppliers']['scope'] = {clean(h): [clean(s), clean(e)] for h, s, e in zip([c.value for c in ws[4]], [c.value for c in ws[2]], [c.value for c in ws[3]]) if clean(h)}
with open(out, 'w') as f: json.dump(data, f, ensure_ascii=False)
print({k: (len(v['rows']) if isinstance(v, dict) and 'rows' in v else (len(v) if hasattr(v, '__len__') else v)) for k, v in data.items() if k not in ('source', 'extractedAt')})
